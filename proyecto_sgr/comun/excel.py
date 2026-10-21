"""
Exportación a Excel de lo que está filtrado en pantalla.

En la maqueta estática esto lo hacía el navegador con una librería externa.
Aquí lo hace el servidor con openpyxl, que es la forma correcta: el archivo
sale del mismo queryset que alimenta la tabla, así que no pueden diferir.

La cabecera del archivo deja escrito qué filtro se exportó. Es lo que
permite, semanas después, saber qué se estaba mirando cuando se generó.
"""

from django.http import HttpResponse
from django.utils import timezone

ANCHO_POR_DEFECTO = 24


def libro_de_listado(*, titulo, filtro, columnas, filas, archivo):
    """Devuelve una respuesta HTTP con el .xlsx del listado."""
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter
    except ImportError:  # pragma: no cover
        return _respaldo_csv(
            titulo=titulo, filtro=filtro, columnas=columnas,
            filas=filas, archivo=archivo,
        )

    ahora = timezone.localtime()
    libro = Workbook()
    hoja = libro.active
    hoja.title = "Listado"

    # El granate oficial del municipio, el mismo de la hoja de estilos.
    granate = "FFAD0000"
    hueso = "FFF1ECE7"

    encabezado = [
        ["Sistema de Gestión de Resultados · Municipalidad de La Serena"],
        [titulo],
        ["Generado", ahora.strftime("%d-%m-%Y %H:%M")],
        ["Filtro aplicado", filtro],
        ["Registros", len(filas)],
        [],
    ]
    for linea in encabezado:
        hoja.append(linea)

    hoja["A1"].font = Font(bold=True, size=12, color=granate)
    hoja["A2"].font = Font(bold=True, size=11)
    for celda in ("A3", "A4", "A5"):
        hoja[celda].font = Font(bold=True, size=9)

    fila_titulos = len(encabezado) + 1
    hoja.append([c.titulo for c in columnas])
    for numero in range(1, len(columnas) + 1):
        celda = hoja.cell(row=fila_titulos, column=numero)
        celda.font = Font(bold=True, size=10)
        celda.fill = PatternFill("solid", fgColor=hueso)
        celda.alignment = Alignment(vertical="center")

    for obj in filas:
        hoja.append([_plano(c.valor(obj)) for c in columnas])

    for numero, columna in enumerate(columnas, start=1):
        ancho = getattr(columna, "ancho", None) or ANCHO_POR_DEFECTO
        hoja.column_dimensions[get_column_letter(numero)].width = ancho

    hoja.freeze_panes = hoja.cell(row=fila_titulos + 1, column=1)

    nombre = f"{archivo}_{ahora.strftime('%Y-%m-%d')}.xlsx"
    respuesta = HttpResponse(
        content_type=(
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )
    )
    respuesta["Content-Disposition"] = f'attachment; filename="{nombre}"'
    libro.save(respuesta)
    return respuesta


def _plano(valor):
    """Deja el valor en algo que Excel entienda."""
    if isinstance(valor, bool):
        return "Sí" if valor else "No"
    if valor is None:
        return ""
    if isinstance(valor, (int, float, str)):
        return valor
    return str(valor)


def csv_de_listado(*, titulo, filtro, columnas, filas, archivo):
    """
    El mismo listado en CSV, para quien lo quiere abrir en otra herramienta.

    Comparte implementación con el respaldo de más abajo: es exactamente el
    mismo archivo, pedido a propósito en vez de por falla. Tenerlo con
    nombre propio evita que alguien lea una llamada a «respaldo» en la
    pantalla de informes y crea que algo salió mal.
    """
    return _respaldo_csv(titulo=titulo, filtro=filtro, columnas=columnas,
                         filas=filas, archivo=archivo)


def _respaldo_csv(*, titulo, filtro, columnas, filas, archivo):
    """
    Si openpyxl no está instalado, se exporta un CSV que Excel abre igual.

    Está por robustez: la exportación es un requisito de la entrega y no
    debería caerse porque falte una dependencia. El separador es punto y
    coma y lleva marca de orden de bytes, que es lo que espera Excel en
    español.
    """
    import csv
    import io

    ahora = timezone.localtime()
    buffer = io.StringIO()
    escritor = csv.writer(buffer, delimiter=";")
    escritor.writerow(["Sistema de Gestión de Resultados · Municipalidad de La Serena"])
    escritor.writerow([titulo])
    escritor.writerow(["Generado", ahora.strftime("%d-%m-%Y %H:%M")])
    escritor.writerow(["Filtro aplicado", filtro])
    escritor.writerow(["Registros", len(filas)])
    escritor.writerow([])
    escritor.writerow([c.titulo for c in columnas])
    for obj in filas:
        escritor.writerow([_plano(c.valor(obj)) for c in columnas])

    nombre = f"{archivo}_{ahora.strftime('%Y-%m-%d')}.csv"
    respuesta = HttpResponse(
        "﻿" + buffer.getvalue(), content_type="text/csv; charset=utf-8"
    )
    respuesta["Content-Disposition"] = f'attachment; filename="{nombre}"'
    return respuesta
